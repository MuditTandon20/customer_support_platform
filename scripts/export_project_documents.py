"""Export the two owner-facing Markdown documents to DOCX and PDF.

Documentation-only dependencies: python-docx and reportlab. Install these in a
separate tooling environment; they are not application runtime requirements.
"""
from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'deliverables'
SOURCES = ['ASTER_PRD_STAGE1', 'ASTER_PROJECT_GUIDE']


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x) for x in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
            continue
        heading = re.match(r'^(#{1,3}) (.+)', line)
        if heading:
            yield 'h' + str(len(heading[1])), heading[2]
            i += 1
            continue
        if line.startswith('- '):
            yield 'bullet', line[2:]
            i += 1
            continue
        parts = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||- )', lines[i].strip()):
            parts.append(lines[i].strip())
            i += 1
        yield 'p', ' '.join(parts)


def prepare_fonts():
    families = [
        (Path('C:/Windows/Fonts/arial.ttf'), Path('C:/Windows/Fonts/arialbd.ttf')),
        (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')),
    ]
    for regular, bold in families:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont('AsterText', str(regular)))
            pdfmetrics.registerFont(TTFont('AsterBold', str(bold)))
            pdfmetrics.registerFontFamily('AsterText', normal='AsterText', bold='AsterBold')
            return
    raise RuntimeError('Install Arial or DejaVu Sans to export Unicode PDF text.')


def write_docx(name, content):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13
    for heading in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[heading].font.name = 'Arial'
        doc.styles[heading].font.color.rgb = RGBColor.from_string('183859')
    header = section.header.paragraphs[0]
    header.text = 'ASTER SUPPORT OS  |  STAGE 1  |  29 OCTOBER 2026'
    header.style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.text = 'Synthetic, non-commercial test release · Not deployed     |     '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'Aster Stage 1 PRD' if name.endswith('STAGE1') else 'Aster — Project explained from scratch'
    doc.core_properties.author = 'Mudit Tandon'
    for kind, value in content:
        if kind.startswith('h'):
            doc.add_heading(value, level=0 if kind == 'h1' else int(kind[1:]) - 1)
        elif kind == 'table':
            table = doc.add_table(rows=1, cols=len(value[0]))
            table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, value[0]):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in value[1:]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
            doc.add_paragraph()
        else:
            doc.add_paragraph(value, style='List Bullet' if kind == 'bullet' else 'Normal')
    if name == 'ASTER_PROJECT_GUIDE':
        doc.add_page_break()
        doc.add_heading('UI preview — local synthetic demo', level=1)
        doc.add_paragraph('Captured during this documentation handoff. Counts describe sample records, not real customer activity.')
        for image_name, caption in [('aster-dashboard.png', 'Overview'), ('aster-inbox.png', 'Inbox')]:
            image_path = OUT / image_name
            if image_path.exists():
                if image_name == 'aster-inbox.png':
                    doc.add_page_break()
                doc.add_paragraph(caption, style='Heading 2')
                doc.add_picture(str(image_path), width=Inches(6.6))
    doc.save(OUT / (name + '.docx'))


def write_pdf(name, content):
    styles = getSampleStyleSheet()
    for style_name in ['Normal', 'Title', 'Heading1', 'Heading2', 'Heading3']:
        s = styles[style_name]
        s.fontName = 'AsterBold' if style_name != 'Normal' else 'AsterText'
        s.textColor = colors.HexColor('#183859') if style_name != 'Normal' else colors.HexColor('#263242')
    styles['Normal'].fontSize = 10
    styles['Normal'].leading = 14.8
    styles['Normal'].spaceAfter = 8
    styles['Title'].fontSize = 25
    styles['Title'].leading = 30
    styles.add(ParagraphStyle(name='Cell', fontName='AsterText', fontSize=8.4, leading=11, spaceAfter=2))
    styles.add(ParagraphStyle(name='SmallCaption', parent=styles['Normal'], fontSize=9, alignment=TA_CENTER))
    flow = []
    width = 487
    for kind, value in content:
        if kind.startswith('h'):
            flow.append(Paragraph(escape(value), styles[{'h1':'Title', 'h2':'Heading1', 'h3':'Heading2'}[kind]]))
        elif kind == 'table':
            rows = [[Paragraph(escape(cell), styles['Cell']) for cell in row] for row in value]
            n = len(value[0])
            widths = [width / n] * n
            if n == 4:
                widths = [42, 130, 207, 108]
            table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E3EDF7')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('LEFTPADDING', (0,0), (-1,-1), 6),
                ('RIGHTPADDING', (0,0), (-1,-1), 6),
                ('TOPPADDING', (0,0), (-1,-1), 5),
                ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ('LINEBELOW', (0,0), (-1,0), .5, colors.HexColor('#B9CADC')),
            ]))
            flow.extend([table, Spacer(1, 10)])
        else:
            flow.append(Paragraph(('• ' if kind == 'bullet' else '') + escape(value), styles['Normal']))
    if name == 'ASTER_PROJECT_GUIDE':
        flow.extend([PageBreak(), Paragraph('UI preview — local synthetic demo', styles['Heading1']),
            Paragraph('Captured during this documentation handoff. These are invented test records.', styles['Normal'])])
        for image_name, caption in [('aster-dashboard.png', 'Overview'), ('aster-inbox.png', 'Inbox')]:
            path = OUT / image_name
            if path.exists():
                if image_name == 'aster-inbox.png':
                    flow.append(PageBreak())
                w, h = ImageReader(str(path)).getSize()
                flow.extend([Paragraph(caption, styles['Heading2']), Image(str(path), width=width, height=width*h/w), Spacer(1, 12)])

    def page(canvas, document):
        canvas.saveState()
        canvas.setFont('AsterText', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(54, 812, 'ASTER SUPPORT OS  |  STAGE 1  |  29 OCTOBER 2026')
        canvas.drawString(54, 29, 'Synthetic, non-commercial test release · Not deployed')
        canvas.drawRightString(541, 29, str(document.page))
        canvas.restoreState()

    SimpleDocTemplate(str(OUT / (name + '.pdf')), pagesize=(595,842),
        rightMargin=54, leftMargin=54, topMargin=52, bottomMargin=48,
        title='Aster Stage 1 PRD' if name.endswith('STAGE1') else 'Aster project guide',
        author='Mudit Tandon').build(flow, onFirstPage=page, onLaterPages=page)


def main():
    OUT.mkdir(exist_ok=True)
    prepare_fonts()
    for name in SOURCES:
        content = list(blocks((ROOT / 'docs' / (name + '.md')).read_text(encoding='utf-8')))
        write_docx(name, content)
        write_pdf(name, content)
        print('Exported', name, 'as DOCX and PDF')


if __name__ == '__main__':
    main()
